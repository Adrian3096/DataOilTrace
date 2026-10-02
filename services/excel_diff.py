from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
MAX_REPORTED_CHANGES = 500


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _workbook_cells(path: str | Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    cells: dict[str, dict[str, Any]] = {}
    try:
        for worksheet in workbook.worksheets:
            sheet_cells: dict[str, Any] = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        sheet_cells[cell.coordinate] = cell.value
            cells[worksheet.title] = sheet_cells
    finally:
        workbook.close()
    return cells


def compare_workbooks(old_path: str | Path, new_path: str | Path) -> dict[str, Any]:
    """Compare workbook cell values and report exact differences and row counts."""
    old_sheets = _workbook_cells(old_path)
    new_sheets = _workbook_cells(new_path)
    changes: list[dict[str, str]] = []
    added_rows = 0
    deleted_rows = 0
    modified_rows = 0
    changed_cell_count = 0
    modified_columns: set[tuple[str, int]] = set()

    for sheet_name in sorted(old_sheets.keys() | new_sheets.keys()):
        old_cells = old_sheets.get(sheet_name, {})
        new_cells = new_sheets.get(sheet_name, {})
        old_row_values: dict[int, dict[int, Any]] = {}
        new_row_values: dict[int, dict[int, Any]] = {}
        for address, value in old_cells.items():
            cell = _split_address(address)
            old_row_values.setdefault(cell[0], {})[cell[1]] = value
        for address, value in new_cells.items():
            cell = _split_address(address)
            new_row_values.setdefault(cell[0], {})[cell[1]] = value

        old_rows, new_rows = set(old_row_values), set(new_row_values)
        added_rows += len(new_rows - old_rows)
        deleted_rows += len(old_rows - new_rows)
        sheet_modified_rows: set[int] = set()

        for address in sorted(old_cells.keys() | new_cells.keys(), key=_address_sort_key):
            old_value = old_cells.get(address)
            new_value = new_cells.get(address)
            if old_value == new_value:
                continue
            changed_cell_count += 1
            row_number, column_number = _split_address(address)
            if row_number in old_rows & new_rows:
                sheet_modified_rows.add(row_number)
                modified_columns.add((sheet_name, column_number))
            if len(changes) < MAX_REPORTED_CHANGES:
                changes.append(
                    {
                        "sheet": sheet_name,
                        "cell": address,
                        "before": _display_value(old_value),
                        "after": _display_value(new_value),
                        "change_type": _cell_change_type(old_value, new_value),
                    }
                )
        modified_rows += len(sheet_modified_rows)

    modified_label = "registro modificado" if modified_rows == 1 else "registros modificados"
    added_label = "registro agregado" if added_rows == 1 else "registros agregados"
    deleted_label = "registro eliminado" if deleted_rows == 1 else "registros eliminados"
    return {
        "modified_rows": modified_rows,
        "added_rows": added_rows,
        "deleted_rows": deleted_rows,
        "modified_columns": len(modified_columns),
        "changed_cells": changed_cell_count,
        "changes": changes,
        "truncated": changed_cell_count > MAX_REPORTED_CHANGES,
        "summary": (
            f"{modified_rows} {modified_label}, {added_rows} {added_label}, "
            f"{deleted_rows} {deleted_label}; {changed_cell_count} celdas cambiaron."
        ),
    }


def _split_address(address: str) -> tuple[int, int]:
    column = "".join(character for character in address if character.isalpha())
    row = "".join(character for character in address if character.isdigit())
    return int(row), _column_number(column)


def _column_number(column: str) -> int:
    number = 0
    for character in column.upper():
        number = number * 26 + ord(character) - ord("A") + 1
    return number


def _address_sort_key(address: str) -> tuple[int, int]:
    row, column = _split_address(address)
    return row, column


def _cell_change_type(old_value: Any, new_value: Any) -> str:
    if old_value is None:
        return "added"
    if new_value is None:
        return "deleted"
    return "modified"
