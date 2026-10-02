import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import storage.version_store as version_store
from services.excel_diff import compare_workbooks
from services.hasher import generate_file_sha256


class VersionHistoryTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.storage_patches = [
            patch.object(version_store, "STORAGE_DIR", self.root / "storage"),
            patch.object(version_store, "SNAPSHOT_DIR", self.root / "storage" / "snapshots"),
            patch.object(version_store, "DATABASE_PATH", self.root / "storage" / "versions.sqlite3"),
        ]
        for storage_patch in self.storage_patches:
            storage_patch.start()
            self.addCleanup(storage_patch.stop)
        version_store.initialize_store()

    def tearDown(self):
        self.temp_dir.cleanup()

    def _make_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        if sheet is None:
            raise AssertionError("openpyxl no creó la hoja inicial")
        sheet.title = "Produccion"
        sheet.append(["Campo", "Barriles"])
        sheet.append(["Norte", 100])
        workbook.save(path)

    def test_diff_reports_edited_cell_and_added_record(self):
        first = self.root / "datos.xlsx"
        second = self.root / "datos_v2.xlsx"
        self._make_workbook(first)
        workbook = load_workbook(first)
        sheet = workbook["Produccion"]
        sheet["B2"] = 125
        sheet.append(["Sur", 50])
        workbook.save(second)

        diff = compare_workbooks(first, second)

        self.assertEqual(diff["modified_rows"], 1)
        self.assertEqual(diff["added_rows"], 1)
        self.assertEqual(diff["modified_columns"], 1)
        self.assertEqual(diff["changed_cells"], 3)
        self.assertEqual(
            [(change["cell"], change["before"], change["after"]) for change in diff["changes"]],
            [("B2", "100", "125"), ("A3", "", "Sur"), ("B3", "", "50")],
        )

    def test_versions_are_saved_with_snapshots_and_case_insensitive_name(self):
        first = self.root / "datos.xlsx"
        second = self.root / "datos_v2.xlsx"
        self._make_workbook(first)
        workbook = load_workbook(first)
        workbook["Produccion"]["B2"] = 125
        workbook.save(second)

        for version, path in enumerate((first, second), start=1):
            file_hash = generate_file_sha256(str(path))
            version_store.register_version(
                file_path=path,
                file_name="datos.xlsx",
                file_hash=file_hash,
                file_size=path.stat().st_size,
                file_type=".xlsx",
                hash_generated_at=f"2026-10-0{version}T10:42:13+00:00",
                topic_id="0.0.1",
                transaction_id=f"tx-{version}",
                sequence_number=version,
                hashscan_url=f"https://hashscan.io/testnet/transaction/tx-{version}",
                diff={"summary": "Versión de prueba", "changes": []},
            )

        records = version_store.get_versions("DATOS.XLSX")
        self.assertEqual([record["version"] for record in records], [1, 2])
        self.assertTrue(all(Path(record["snapshot_path"]).is_file() for record in records))
        self.assertEqual(Path(records[0]["source_path"]), first.resolve())
        self.assertEqual(records[0]["hash_generated_at"], "2026-10-01T10:42:13+00:00")
        self.assertEqual(records[1]["hash_generated_at"], "2026-10-02T10:42:13+00:00")
        latest = version_store.get_latest("datos.xlsx")
        self.assertIsNotNone(latest)
        assert latest is not None
        self.assertEqual(latest["version"], 2)
        first_version = version_store.get_version("datos.xlsx", 1)
        self.assertIsNotNone(first_version)
        assert first_version is not None
        self.assertEqual(first_version["version"], 1)

    def test_dashboard_aggregates_distinct_files_versions_and_latest_paths(self):
        first = self.root / "produccion.xlsx"
        another = self.root / "pozos.xlsx"
        self._make_workbook(first)
        self._make_workbook(another)
        for file_path, file_name, tx in (
            (first, "produccion.xlsx", "tx-p1"),
            (another, "pozos.xlsx", "tx-z1"),
        ):
            version_store.register_version(
                file_path=file_path,
                file_name=file_name,
                file_hash=generate_file_sha256(str(file_path)),
                file_size=file_path.stat().st_size,
                file_type=".xlsx",
                topic_id="0.0.1",
                transaction_id=tx,
                sequence_number=1,
                hashscan_url="https://example.test/tx",
                diff={"summary": "Versión inicial", "changes": []},
            )

        workbook = load_workbook(first)
        workbook["Produccion"]["B2"] = 200
        workbook.save(first)
        version_store.register_version(
            file_path=first,
            file_name="produccion.xlsx",
            file_hash=generate_file_sha256(str(first)),
            file_size=first.stat().st_size,
            file_type=".xlsx",
            topic_id="0.0.1",
            transaction_id="tx-p2",
            sequence_number=2,
            hashscan_url="https://example.test/tx-p2",
            diff={"summary": "Cambio de prueba", "changes": []},
        )

        data = version_store.get_dashboard_data(limit=10)

        self.assertEqual(data["file_count"], 2)
        self.assertEqual(data["version_count"], 3)
        self.assertEqual(data["record_count"], 3)
        self.assertEqual(len(data["files"]), 2)
        latest = next(item for item in data["files"] if item["file_name"] == "produccion.xlsx")
        self.assertEqual(latest["version"], 2)
        self.assertEqual(latest["version_count"], 2)
        version_store.update_source_path("produccion.xlsx", another)
        updated = version_store.get_latest("produccion.xlsx")
        self.assertIsNotNone(updated)
        assert updated is not None
        self.assertEqual(Path(updated["source_path"]), another.resolve())

    def test_snapshot_is_not_saved_if_source_hash_changed(self):
        file_path = self.root / "datos.xlsx"
        self._make_workbook(file_path)

        with self.assertRaises(ValueError):
            version_store.register_version(
                file_path=file_path,
                file_name=file_path.name,
                file_hash="not-the-current-hash",
                file_size=file_path.stat().st_size,
                file_type=".xlsx",
                topic_id="0.0.1",
                transaction_id="tx-invalid",
                sequence_number=1,
                hashscan_url="https://example.test",
                diff={"summary": "invalid", "changes": []},
            )
        self.assertEqual(version_store.get_versions(file_path.name), [])

    def test_initialize_store_migrates_existing_version_table(self):
        legacy_db = self.root / "legacy.sqlite3"
        with patch.object(version_store, "DATABASE_PATH", legacy_db):
            legacy_db.parent.mkdir(parents=True, exist_ok=True)
            connection = sqlite3.connect(legacy_db)
            try:
                connection.execute(
                    """
                    CREATE TABLE file_versions (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        file_key TEXT NOT NULL,
                        file_name TEXT NOT NULL,
                        version INTEGER NOT NULL,
                        file_hash TEXT NOT NULL,
                        file_size INTEGER NOT NULL,
                        file_type TEXT NOT NULL,
                        registered_at TEXT NOT NULL,
                        topic_id TEXT NOT NULL,
                        transaction_id TEXT NOT NULL,
                        sequence_number INTEGER,
                        hashscan_url TEXT NOT NULL,
                        snapshot_path TEXT NOT NULL,
                        diff_json TEXT NOT NULL
                    )
                    """
                )
                connection.commit()
            finally:
                connection.close()

            version_store.initialize_store()
            connection = sqlite3.connect(legacy_db)
            try:
                columns = {
                    row[1]
                    for row in connection.execute("PRAGMA table_info(file_versions)")
                }
            finally:
                connection.close()
        self.assertIn("hash_generated_at", columns)
        self.assertIn("source_path", columns)


if __name__ == "__main__":
    unittest.main()
